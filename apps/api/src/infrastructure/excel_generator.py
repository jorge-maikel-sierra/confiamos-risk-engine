"""Generador de planillas de trabajo financiero en Excel (.xlsx) (SPEC-07)."""

import io
from typing import Any

import openpyxl
from openpyxl.styles import Border, Font, PatternFill, Side


class ExcelFinancialGenerator:
    """Construye un libro de cálculo Excel (.xlsx) con fórmulas financieras vivas y auditoría."""

    def __init__(self) -> None:
        self.header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        self.header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        self.title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        self.bold_font = Font(name="Calibri", size=11, bold=True)
        self.regular_font = Font(name="Calibri", size=11)
        self.hash_font = Font(name="Courier New", size=9, color="475569")
        self.thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

    def generate(self, data: dict[str, Any]) -> bytes:
        """Genera el contenido binario del archivo Excel (.xlsx)."""
        wb = openpyxl.Workbook()

        # 1. Hoja 1: Resumen Ejecutivo
        ws_summary = wb.active
        ws_summary.title = "Resumen Ejecutivo"
        self._build_summary_sheet(ws_summary, data)

        # 2. Hoja 2: Balance y Liquidez
        ws_liquidity = wb.create_sheet(title="Balance y Liquidez")
        self._build_liquidity_sheet(ws_liquidity, data)

        # 3. Hoja 3: Flujo y Cobertura
        ws_cashflow = wb.create_sheet(title="Flujo y Cobertura")
        self._build_cashflow_sheet(ws_cashflow, data)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def _build_summary_sheet(self, ws: Any, data: dict[str, Any]) -> None:
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 50

        ws["A1"] = "COOPERATIVA CONFIAMOS"
        ws["A1"].font = self.title_font
        ws["A2"] = "DICTAMEN FINANCIERO Y MODELO DE SUSCRIPCIÓN"
        ws["A2"].font = Font(name="Calibri", size=11, bold=True, color="2563EB")

        rows = [
            ("Folio de Solicitud:", str(data.get("id", "N/A"))),
            ("Razón Social:", str(data.get("business_name", "N/A"))),
            ("Identificación / NIT:", str(data.get("associated_id", "N/A"))),
            ("Monto Solicitado (USD):", float(data.get("requested_amount_usd", 0.0))),
            ("Plazo (Meses):", int(data.get("term_months", 0))),
            ("Fecha de Radicación:", str(data.get("created_at", "N/A"))),
            ("Recomendación IA:", str(data.get("ai_verdict", {}).get("recommendation", "N/A"))),
            ("Sello de Integridad SHA-256:", str(data.get("crypto_hash", "NO_HASH"))),
        ]

        curr_row = 4
        for label, val in rows:
            ws.cell(row=curr_row, column=1, value=label).font = self.bold_font
            cell_val = ws.cell(row=curr_row, column=2, value=val)
            if label == "Monto Solicitado (USD):":
                cell_val.number_format = "$#,##0.00"
                cell_val.font = self.regular_font
            elif "SHA-256" in label:
                cell_val.font = self.hash_font
            else:
                cell_val.font = self.regular_font
            curr_row += 1

    def _build_liquidity_sheet(self, ws: Any, data: dict[str, Any]) -> None:
        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 25

        ws["A1"] = "ANÁLISIS DE BALANCE Y LIQUIDEZ CORRIENTE"
        ws["A1"].font = self.title_font

        ws["A3"] = "Concepto Contable"
        ws["B3"] = "Valor (USD)"
        ws["C3"] = "Referencia / Fórmula"
        for col in ["A3", "B3", "C3"]:
            ws[col].fill = self.header_fill
            ws[col].font = self.header_font

        raw = data.get("raw_financials") or {}
        # Si no viene en raw_financials, deducir de métricas o defaults
        ca_est = data.get("metrics", {}).get("current_ratio", 1.5) * 100000.0
        ca = float(raw.get("current_assets") or ca_est)
        cl = float(raw.get("current_liabilities") or 100000.0)

        ws["A4"] = "Activo Corriente Total"
        ws["B4"] = ca
        ws["B4"].number_format = "$#,##0.00"
        ws["C4"] = "Ingesta balances contables"

        ws["A5"] = "Pasivo Corriente Total"
        ws["B5"] = cl
        ws["B5"].number_format = "$#,##0.00"
        ws["C5"] = "Deuda exigible a corto plazo"

        ws["A6"] = "Ratio de Liquidez Corriente"
        ws["B6"] = "=B4/B5"
        ws["B6"].number_format = "0.00x"
        ws["B6"].font = self.bold_font
        ws["C6"] = "Benchmark: ≥ 1.30x"

        for row in range(4, 7):
            for col in range(1, 4):
                ws.cell(row=row, column=col).border = self.thin_border

    def _build_cashflow_sheet(self, ws: Any, data: dict[str, Any]) -> None:
        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 28

        ws["A1"] = "ANÁLISIS DE FLUJO DE CAJA Y COBERTURA DE DEUDA"
        ws["A1"].font = self.title_font

        ws["A3"] = "Métrica Financiera"
        ws["B3"] = "Monto (USD)"
        ws["C3"] = "Regla de Suscripción"
        for col in ["A3", "B3", "C3"]:
            ws[col].fill = self.header_fill
            ws[col].font = self.header_font

        metrics = data.get("metrics") or {}
        ocf = float(metrics.get("operating_cash_flow", 0.0))
        ads = float(metrics.get("annual_debt_service", 0.0))

        ws["A4"] = "Flujo de Caja Operativo Anual (OCF)"
        ws["B4"] = ocf
        ws["B4"].number_format = "$#,##0.00"
        ws["C4"] = "Capacidad neta generadora"

        ws["A5"] = "Servicio Anual de la Deuda (ADS)"
        ws["B5"] = ads
        ws["B5"].number_format = "$#,##0.00"
        ws["C5"] = "Amortizaciones + Intereses"

        ws["A6"] = "DSCR (Debt Service Coverage Ratio)"
        ws["B6"] = '=IF(B5=0,"N/A",B4/B5)'
        ws["B6"].number_format = "0.00x"
        ws["B6"].font = self.bold_font
        ws["C6"] = "Mínimo exigido: ≥ 1.25x"

        for row in range(4, 7):
            for col in range(1, 4):
                ws.cell(row=row, column=col).border = self.thin_border
